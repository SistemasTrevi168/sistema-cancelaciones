"""
KPIs del panel de administración: rendimiento de cada usuario u operador.

- registrar(): guarda en la tabla "actividad" lo que hace cada usuario (login, leer, guardar, word, zip...).
  Usa su propia conexión y nunca lanza errores: si falla, el sistema sigue funcionando igual.
- /api/admin/kpis: resumen, producción por día / hora / día de la semana, tabla por operador,
  plantillas más usadas y campos que más corrigen a mano.
- /api/admin/kpis/excel: lo mismo en un archivo de Excel.

La producción se calcula con la tabla de expedientes (funciona también con los expedientes
de antes). Los tiempos de revisión y las correcciones solo existen desde que se agregó la bitácora.
"""
import io
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, time
from typing import Any, Dict, List, Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

import models
from auth import require_admin
from database import SessionLocal, get_db

ZONA = ZoneInfo("America/Mexico_City")
DIAS_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
ORIGENES = {"individual": "Caso individual", "masivo": "Carga masiva", "manual": "Captura manual",
            "excel": "Carga Excel", "anterior": "Sin dato (antes de la bitácora)"}
# Más de 4 horas entre leer y guardar = el borrador se quedó abierto; no cuenta para el tiempo
MAX_REVISION_SEG = 4 * 3600

router = APIRouter(prefix="/api/admin/kpis", tags=["kpis"])


# --------------------------------------------------------------------------------------
# Registro de actividad
# --------------------------------------------------------------------------------------
def registrar(username: str, tipo: str, expediente_id: Any = None, numero_credito: Any = None, **detalle):
    try:
        db = SessionLocal()
        try:
            db.add(models.Actividad(
                username=str(username or "desconocido"),
                tipo=tipo,
                expediente_id=str(expediente_id) if expediente_id else None,
                numero_credito=str(numero_credito) if numero_credito else None,
                detalle=detalle or None,
            ))
            db.commit()
        finally:
            db.close()
    except Exception as e:  # nunca debe romper la operación principal
        print(f"Aviso: no se pudo registrar la actividad '{tipo}': {e}")


def metricas_de_guardado(metricas: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Convierte lo que manda el frontend al guardar en datos para la bitácora."""
    metricas = metricas or {}
    salida: Dict[str, Any] = {"origen": metricas.get("origen") or "individual"}
    leido_en = metricas.get("leido_en")
    try:
        if leido_en:
            seg = (datetime.now().timestamp() * 1000 - float(leido_en)) / 1000
            if 0 <= seg <= MAX_REVISION_SEG:
                salida["segundos_revision"] = round(seg)
    except (TypeError, ValueError):
        pass
    corregidos = [str(c) for c in (metricas.get("corregidos") or []) if c][:60]
    # Valores antes -> después de lo que se corrigió a mano (para el historial de cambios del expediente)
    cambios = metricas.get("cambios") or {}
    if isinstance(cambios, dict):
        salida["cambios"] = {str(k)[:60]: [str(v[0] if isinstance(v, (list, tuple)) and v else "")[:300],
                                           str(v[1] if isinstance(v, (list, tuple)) and len(v) > 1 else "")[:300]]
                             for k, v in list(cambios.items())[:60]}
    if leido_en:
        salida["leido_en"] = leido_en
    salida["corregidos"] = corregidos
    salida["num_corregidos"] = len(corregidos)
    try:
        salida["avisos"] = int(metricas.get("avisos") or 0)
    except (TypeError, ValueError):
        salida["avisos"] = 0
    return salida


# --------------------------------------------------------------------------------------
# Cálculo
# --------------------------------------------------------------------------------------
def _local(dt: Optional[datetime]) -> Optional[datetime]:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZoneInfo("UTC"))
    return dt.astimezone(ZONA)


def _rango(desde: Optional[str], hasta: Optional[str]):
    hoy = datetime.now(ZONA).date()
    try:
        d_hasta = date.fromisoformat(hasta) if hasta else hoy
        d_desde = date.fromisoformat(desde) if desde else d_hasta - timedelta(days=29)
    except ValueError:
        raise HTTPException(status_code=400, detail="Fechas inválidas (usa AAAA-MM-DD).")
    if d_desde > d_hasta:
        d_desde, d_hasta = d_hasta, d_desde
    if (d_hasta - d_desde).days > 731:
        raise HTTPException(status_code=400, detail="El periodo no puede ser mayor a 2 años.")
    return d_desde, d_hasta


def _limites(d_desde: date, d_hasta: date):
    inicio = datetime.combine(d_desde, time.min, tzinfo=ZONA)
    fin = datetime.combine(d_hasta + timedelta(days=1), time.min, tzinfo=ZONA)
    return inicio, fin


def _hhmm(minutos: Optional[float]) -> str:
    if minutos is None:
        return ""
    minutos = int(round(minutos))
    return f"{minutos // 60:02d}:{minutos % 60:02d}"


def _expedientes(db: Session, inicio, fin, usuario: Optional[str]):
    E = models.Expediente
    q = db.query(E.id, E.usuario_propietario, E.numero_credito, E.fecha_creacion, E.estado,
                 E.datos_extraidos.op("->>")("plantilla_seleccionada")) \
        .filter(E.fecha_creacion >= inicio, E.fecha_creacion < fin)
    if usuario:
        q = q.filter(E.usuario_propietario == usuario)
    return [{"id": str(r[0]), "usuario": r[1], "credito": r[2], "fecha": _local(r[3]), "estado": r[4],
             "plantilla": (r[5] or "").replace(".docx", "")} for r in q.all()]


def _actividad(db: Session, inicio, fin, usuario: Optional[str]):
    A = models.Actividad
    # Las correcciones internas del sistema no son trabajo de nadie: no cuentan en los reportes
    q = db.query(A).filter(A.creado_en >= inicio, A.creado_en < fin,
                           A.tipo.notin_(("corregir_plantilla", "mantenimiento", "documentos")))
    if usuario:
        q = q.filter(A.username == usuario)
    return [{"usuario": a.username, "tipo": a.tipo, "expediente_id": a.expediente_id,
             "credito": a.numero_credito, "detalle": a.detalle or {}, "fecha": _local(a.creado_en)} for a in q.all()]


def _mediana(valores):
    return round(statistics.median(valores)) if valores else None


def calcular(db: Session, desde: Optional[str], hasta: Optional[str], usuario: Optional[str]) -> Dict[str, Any]:
    d_desde, d_hasta = _rango(desde, hasta)
    inicio, fin = _limites(d_desde, d_hasta)
    usuario = (usuario or "").strip() or None

    exps = _expedientes(db, inicio, fin, usuario)
    acts = _actividad(db, inicio, fin, usuario)

    # Periodo anterior del mismo largo (para comparar)
    dias = (d_hasta - d_desde).days + 1
    ant_inicio, ant_fin = _limites(d_desde - timedelta(days=dias), d_desde - timedelta(days=1))
    E = models.Expediente
    q_ant = db.query(E.id).filter(E.fecha_creacion >= ant_inicio, E.fecha_creacion < ant_fin)
    if usuario:
        q_ant = q_ant.filter(E.usuario_propietario == usuario)
    total_anterior = q_ant.count()

    guardados = {a["expediente_id"]: a for a in acts if a["tipo"] in ("guardar", "manual", "excel") and a["expediente_id"]}

    def origen_de(exp):
        a = guardados.get(exp["id"])
        if not a:
            return "anterior"
        if a["tipo"] in ("manual", "excel"):
            return a["tipo"]
        return a["detalle"].get("origen") or "individual"

    for e in exps:
        e["origen"] = origen_de(e)
        det = (guardados.get(e["id"]) or {}).get("detalle") or {}
        e["segundos_revision"] = det.get("segundos_revision")
        e["corregidos"] = det.get("corregidos") or []
        e["avisos"] = det.get("avisos") or 0

    # ---------------- Por operador ----------------
    nombres = sorted({e["usuario"] for e in exps} | {a["usuario"] for a in acts})
    por_usuario = []
    for nombre in nombres:
        mis_exps = [e for e in exps if e["usuario"] == nombre]
        mis_acts = [a for a in acts if a["usuario"] == nombre]
        momentos = [e["fecha"] for e in mis_exps] + [a["fecha"] for a in mis_acts]
        por_dia: Dict[date, List[datetime]] = defaultdict(list)
        for m in momentos:
            por_dia[m.date()].append(m)
        inicios = [min(v).hour * 60 + min(v).minute for v in por_dia.values()]
        finales = [max(v).hour * 60 + max(v).minute for v in por_dia.values()]
        jornadas = [(max(v) - min(v)).total_seconds() / 3600 for v in por_dia.values()]
        dias_con_exp = len({e["fecha"].date() for e in mis_exps})
        guard = [a for a in mis_acts if a["tipo"] == "guardar"]
        tiempos = [a["detalle"].get("segundos_revision") for a in guard if a["detalle"].get("segundos_revision") is not None]
        correcciones = [a["detalle"].get("num_corregidos", 0) for a in guard]
        lecturas = [a for a in mis_acts if a["tipo"] == "leer"]
        leidos = sum(int(a["detalle"].get("expedientes") or 0) for a in lecturas)
        zips = [a for a in mis_acts if a["tipo"] == "zip"]
        horas = Counter(e["fecha"].hour for e in mis_exps)
        hora_pico = max(horas.items(), key=lambda x: x[1])[0] if horas else None
        plantillas = Counter(e["plantilla"] for e in mis_exps if e["plantilla"])
        ultima = max(momentos) if momentos else None
        por_usuario.append({
            "usuario": nombre,
            "expedientes": len(mis_exps),
            "dias_activos": len(por_dia),
            "promedio_por_dia": round(len(mis_exps) / dias_con_exp, 1) if dias_con_exp else 0,
            "hora_inicio_prom": _hhmm(statistics.mean(inicios)) if inicios else "",
            "hora_fin_prom": _hhmm(statistics.mean(finales)) if finales else "",
            "jornada_prom_horas": round(statistics.mean(jornadas), 1) if jornadas else 0,
            "hora_pico": f"{hora_pico:02d}:00" if hora_pico is not None else "",
            "revision_mediana_seg": _mediana(tiempos),
            "corregidos_prom": round(statistics.mean(correcciones), 1) if correcciones else None,
            "pct_con_correcciones": round(100 * sum(1 for c in correcciones if c) / len(correcciones)) if correcciones else None,
            "con_avisos": sum(1 for a in guard if a["detalle"].get("avisos")),
            "leidos": leidos,
            "pct_guardados": round(100 * len(guard) / leidos) if leidos else None,
            "words": sum(1 for a in mis_acts if a["tipo"] == "word"),
            "zips": len(zips),
            "expedientes_en_zip": sum(int(a["detalle"].get("expedientes") or 0) for a in zips),
            "sesiones": sum(1 for a in mis_acts if a["tipo"] == "login"),
            "plantilla_top": plantillas.most_common(1)[0][0] if plantillas else "",
            "ultima_actividad": ultima.strftime("%Y-%m-%d %H:%M") if ultima else "",
        })
    por_usuario.sort(key=lambda u: (-u["expedientes"], u["usuario"]))

    # ---------------- Series ----------------
    serie_dia = []
    conteo_dia = Counter(e["fecha"].date() for e in exps)
    conteo_dia_usr: Dict[date, Counter] = defaultdict(Counter)
    for e in exps:
        conteo_dia_usr[e["fecha"].date()][e["usuario"]] += 1
    d = d_desde
    while d <= d_hasta:
        serie_dia.append({"fecha": d.isoformat(), "total": conteo_dia.get(d, 0), "por_usuario": dict(conteo_dia_usr.get(d, {}))})
        d += timedelta(days=1)

    hora_usr: Dict[int, Counter] = defaultdict(Counter)
    sem_usr: Dict[int, Counter] = defaultdict(Counter)
    for e in exps:
        hora_usr[e["fecha"].hour][e["usuario"]] += 1
        sem_usr[e["fecha"].weekday()][e["usuario"]] += 1
    por_hora = [{"hora": h, "total": sum(hora_usr[h].values()), "por_usuario": dict(hora_usr[h])} for h in range(24)]
    por_dia_semana = [{"dia": DIAS_SEMANA[i], "total": sum(sem_usr[i].values()), "por_usuario": dict(sem_usr[i])} for i in range(7)]

    plantillas = Counter(e["plantilla"] or "Sin plantilla" for e in exps)
    campos = Counter(c for a in acts if a["tipo"] == "guardar" for c in (a["detalle"].get("corregidos") or []))
    origenes = Counter(e["origen"] for e in exps)

    todas_guard = [a for a in acts if a["tipo"] == "guardar"]
    tiempos = [a["detalle"].get("segundos_revision") for a in todas_guard if a["detalle"].get("segundos_revision") is not None]
    corr = [a["detalle"].get("num_corregidos", 0) for a in todas_guard]
    dias_con_exp = len(conteo_dia)
    mejor_dia = max(conteo_dia.items(), key=lambda x: x[1]) if conteo_dia else None

    resumen = {
        "expedientes": len(exps),
        "expedientes_periodo_anterior": total_anterior,
        "cambio_pct": round(100 * (len(exps) - total_anterior) / total_anterior) if total_anterior else None,
        "operadores_activos": len([u for u in por_usuario if u["dias_activos"]]),
        "promedio_por_dia": round(len(exps) / dias_con_exp, 1) if dias_con_exp else 0,
        "dias_con_produccion": dias_con_exp,
        "mejor_dia": {"fecha": mejor_dia[0].isoformat(), "total": mejor_dia[1]} if mejor_dia else None,
        "revision_mediana_seg": _mediana(tiempos),
        "pct_con_correcciones": round(100 * sum(1 for c in corr if c) / len(corr)) if corr else None,
        "con_avisos": sum(1 for a in todas_guard if a["detalle"].get("avisos")),
        "words": sum(1 for a in acts if a["tipo"] == "word"),
        "zips": sum(1 for a in acts if a["tipo"] == "zip"),
        "con_bitacora": len(todas_guard) > 0,
    }

    return {
        "desde": d_desde.isoformat(),
        "hasta": d_hasta.isoformat(),
        "usuario": usuario,
        "resumen": resumen,
        "por_usuario": por_usuario,
        "por_dia": serie_dia,
        "por_hora": por_hora,
        "por_dia_semana": por_dia_semana,
        "plantillas": [{"plantilla": k, "total": v} for k, v in plantillas.most_common(12)],
        "campos_corregidos": [{"campo": k, "veces": v} for k, v in campos.most_common(12)],
        "origenes": [{"origen": k, "nombre": ORIGENES.get(k, k), "total": v} for k, v in origenes.most_common()],
        "_expedientes": exps,
    }


# --------------------------------------------------------------------------------------
# Endpoints
# --------------------------------------------------------------------------------------
@router.get("")
def obtener_kpis(
    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),
    usuario: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    admin: dict = Depends(require_admin),
):
    datos = calcular(db, desde, hasta, usuario)
    datos.pop("_expedientes", None)
    datos["usuarios"] = [{"username": u.username, "activo": getattr(u, "activo", True), "nombre": getattr(u, "nombre_mostrar", None) or ""}
                         for u in db.query(models.Usuario).order_by(models.Usuario.username).all()]
    return datos


def _segundos_texto(seg) -> str:
    if seg is None:
        return ""
    seg = int(seg)
    return f"{seg // 60} min {seg % 60:02d} s" if seg >= 60 else f"{seg} s"


@router.get("/excel")
def exportar_kpis_excel(
    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),
    usuario: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    admin: dict = Depends(require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    k = calcular(db, desde, hasta, usuario)
    wb = Workbook()
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="2563EB")

    def hoja(titulo, columnas, filas, primera=False):
        ws = wb.active if primera else wb.create_sheet()
        ws.title = titulo
        ws.append(columnas)
        for c in ws[1]:
            c.font, c.fill = encabezado, relleno
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for f in filas:
            ws.append(list(f))
        for i, col in enumerate(columnas, start=1):
            largo = max([len(str(col))] + [len(str(f[i - 1])) for f in filas if i - 1 < len(f) and f[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)
        ws.freeze_panes = "A2"
        return ws

    r = k["resumen"]
    periodo = f'{k["desde"]} a {k["hasta"]}' + (f' — {k["usuario"]}' if k["usuario"] else " — todos los usuarios")
    hoja("Resumen", ["Indicador", "Valor"], [
        ["Periodo", periodo],
        ["Expedientes generados", r["expedientes"]],
        ["Periodo anterior (mismo largo)", r["expedientes_periodo_anterior"]],
        ["Cambio vs periodo anterior", f'{r["cambio_pct"]}%' if r["cambio_pct"] is not None else ""],
        ["Operadores activos", r["operadores_activos"]],
        ["Promedio por día con producción", r["promedio_por_dia"]],
        ["Días con producción", r["dias_con_produccion"]],
        ["Mejor día", f'{r["mejor_dia"]["fecha"]} ({r["mejor_dia"]["total"]})' if r["mejor_dia"] else ""],
        ["Tiempo típico de revisión (mediana)", _segundos_texto(r["revision_mediana_seg"])],
        ["% guardados con correcciones a mano", f'{r["pct_con_correcciones"]}%' if r["pct_con_correcciones"] is not None else ""],
        ["Expedientes guardados con avisos", r["con_avisos"]],
        ["Word descargados", r["words"]],
        ["ZIP descargados", r["zips"]],
    ], primera=True)

    hoja("Por operador", [
        "Operador", "Expedientes", "Días activos", "Promedio por día", "Hora de inicio (prom.)", "Hora de fin (prom.)",
        "Jornada prom. (h)", "Hora pico", "Revisión típica", "Campos corregidos (prom.)", "% con correcciones",
        "Con avisos", "Leídos", "% leídos que se guardaron", "Word", "ZIP", "Expedientes en ZIP", "Inicios de sesión",
        "Plantilla más usada", "Última actividad"], [
        [u["usuario"], u["expedientes"], u["dias_activos"], u["promedio_por_dia"], u["hora_inicio_prom"], u["hora_fin_prom"],
         u["jornada_prom_horas"], u["hora_pico"], _segundos_texto(u["revision_mediana_seg"]), u["corregidos_prom"],
         u["pct_con_correcciones"], u["con_avisos"], u["leidos"], u["pct_guardados"], u["words"], u["zips"],
         u["expedientes_en_zip"], u["sesiones"], u["plantilla_top"], u["ultima_actividad"]] for u in k["por_usuario"]])

    operadores = [u["usuario"] for u in k["por_usuario"]]
    hoja("Por día", ["Fecha", "Día", "Total"] + operadores, [
        [d["fecha"], DIAS_SEMANA[date.fromisoformat(d["fecha"]).weekday()], d["total"]] + [d["por_usuario"].get(o, 0) for o in operadores]
        for d in k["por_dia"]])
    hoja("Por hora", ["Hora", "Expedientes"], [[f'{h["hora"]:02d}:00 - {h["hora"]:02d}:59', h["total"]] for h in k["por_hora"]])
    hoja("Por día de la semana", ["Día", "Expedientes"], [[d["dia"], d["total"]] for d in k["por_dia_semana"]])
    hoja("Plantillas", ["Plantilla", "Expedientes"], [[p["plantilla"], p["total"]] for p in k["plantillas"]])
    hoja("Campos corregidos", ["Campo", "Veces que se corrigió a mano"], [[c["campo"], c["veces"]] for c in k["campos_corregidos"]])
    hoja("Expedientes", ["Fecha", "Hora", "Operador", "Número de crédito", "Plantilla", "Origen", "Revisión",
                         "Campos corregidos", "Avisos"], [
        [e["fecha"].strftime("%Y-%m-%d"), e["fecha"].strftime("%H:%M"), e["usuario"], e["credito"], e["plantilla"],
         ORIGENES.get(e["origen"], e["origen"]), _segundos_texto(e["segundos_revision"]), ", ".join(e["corregidos"]), e["avisos"]]
        for e in sorted(k["_expedientes"], key=lambda x: x["fecha"])])

    salida = io.BytesIO()
    wb.save(salida)
    nombre = f'KPIs_{k["desde"]}_a_{k["hasta"]}' + (f'_{k["usuario"]}' if k["usuario"] else "") + ".xlsx"
    return Response(content=salida.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"'})
