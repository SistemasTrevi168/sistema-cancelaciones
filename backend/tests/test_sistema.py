"""
Pruebas automáticas del backend (las corre GitHub en cada push, y se pueden correr a mano).

    cd backend
    DATABASE_URL=postgresql+psycopg://usuario@localhost:5432/base python tests/test_sistema.py

No usan PDFs reales (tienen datos personales): trabajan con datos de ejemplo.
Regresan código 1 si alguna prueba falla.
"""
import io
import json
import os
import shutil
import sys
import uuid
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("RESPALDOS_DESACTIVADOS", "1")

import docx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import main  # noqa: E402

FALLAS = []


def chk(nombre, condicion, extra=""):
    print(("OK    " if condicion else "FALLA ") + nombre + (f"  -> {extra}" if (extra and not condicion) else ""))
    if not condicion:
        FALLAS.append(nombre)


def texto_word(contenido: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(contenido)).paragraphs)


def cliente():
    return TestClient(main.app, base_url="https://testserver")


SUF = uuid.uuid4().hex[:6]
CREDITO = "15" + str(uuid.uuid4().int)[:8]
DATOS = {
    "nombre_acreditado": "MARIA PRUEBA EJEMPLO", "numero_credito": CREDITO, "monto_credito": "$150,000.00",
    "numero_carta": "0543910959", "entidad_financiera": "INFONAVIT", "fecha_liquidacion": "20/06/2024",
    "folio_real": "37509", "oficina_registral": "Toluca", "fecha_expedicion": "11/08/2026",
    "datos_inmueble": "Vivienda: DEPARTAMENTO 2, Lote: 2, Manzana: 10, Municipio: METEPEC, Estado: ESTADO DE MÉXICO",
    "credito_a_salario": "45.2095", "numero_escritura": "6411", "fecha_escritura": "10/09/2018",
    "notario_origen_completo": "del Licenciado Jorge Valdes Ramirez, Notario Público número veinticuatro del Estado de México",
}
PLANTILLA = "EDOMEX_AP_M_SOLTERA"

with cliente() as admin, cliente() as ana, cliente() as beto:
    r = admin.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")})
    chk("login admin", r.status_code == 200, r.text[:200])
    for nombre in ("ana", "beto"):
        r = admin.post("/api/admin/usuarios", data={"username": f"{nombre}_{SUF}", "password": "Prueba2026!", "es_admin": "false"})
        chk(f"crear usuario {nombre}", r.status_code == 200, r.text[:200])
    chk("login ana", ana.post("/api/auth/login", json={"username": f"ana_{SUF}", "password": "Prueba2026!"}).status_code == 200)
    chk("login beto", beto.post("/api/auth/login", json={"username": f"beto_{SUF}", "password": "Prueba2026!"}).status_code == 200)

    # ---- Permisos básicos
    chk("sin sesión no entra", cliente().get("/api/expedientes").status_code == 401)
    chk("operador no ve usuarios", ana.get("/api/admin/usuarios").status_code == 403)
    chk("operador no ve respaldos", ana.get("/api/admin/respaldos").status_code == 403)

    # ---- Créditos repetidos (antes de guardar no hay)
    r = ana.get("/api/expedientes/repetidos", params={"creditos": CREDITO}).json()
    chk("crédito nuevo no aparece como repetido", CREDITO not in r["repetidos"])

    # ---- Guardar (lee -> corrige -> guarda): queda GENERADO, con versión de plantilla y bitácora
    metricas = {"origen": "individual", "corregidos": ["fecha_escritura"], "cambios": {"fecha_escritura": ["", "10/09/2018"]}}
    r = ana.post("/api/expedientes/guardar", json={"expedientes": [{"ref": "x", "datos": DATOS, "plantilla": PLANTILLA,
                                                                    "archivos": [f"{CREDITO}_Carta.pdf"], "metricas": metricas}]})
    res = r.json()["resultados"][0]
    chk("guardar expediente", r.status_code == 200 and res.get("expediente_id"), r.text[:300])
    exp_id = res.get("expediente_id")
    chk("guarda la versión de plantilla", bool(res["datos_extraidos"].get("_plantilla_version")))

    hist = {h["id"]: h for h in ana.get("/api/expedientes").json()}
    chk("aparece en Mi Historial como GENERADO", hist.get(exp_id, {}).get("estado") == "GENERADO")
    chk("la plantilla no ha cambiado", hist.get(exp_id, {}).get("plantilla_cambio") is False)
    chk("beto no ve el expediente de ana", exp_id not in {h["id"] for h in beto.get("/api/expedientes").json()})

    r = beto.get("/api/expedientes/repetidos", params={"creditos": f"{CREDITO},123"}).json()["repetidos"]
    chk("repetido: avisa quién y cuándo (aunque sea de otro usuario)",
        CREDITO in r and r[CREDITO][0]["usuario"] == f"ana_{SUF}" and r[CREDITO][0]["fecha"])

    # ---- Estados
    r = ana.put(f"/api/expedientes/{exp_id}/estado", json={"estado": "revisado"})
    chk("cambiar a REVISADO", r.status_code == 200 and r.json()["estado"] == "REVISADO", r.text[:200])
    chk("estado inválido -> 400", ana.put(f"/api/expedientes/{exp_id}/estado", json={"estado": "FIRMADO"}).status_code == 400)
    chk("beto no cambia el estado de ana", beto.put(f"/api/expedientes/{exp_id}/estado", json={"estado": "ENTREGADO"}).status_code == 404)
    r = ana.post("/api/expedientes/estado", json={"ids": [exp_id], "estado": "ENTREGADO"})
    chk("cambio de estado de varios", r.json().get("cambiados") == 1)

    # ---- Editar y volver a descargar: queda en la bitácora con antes -> después
    datos2 = dict(DATOS, nombre_acreditado="MARIA PRUEBA CORREGIDA")
    r = ana.post(f"/api/expedientes/{exp_id}/generar-word", json={"datos": datos2, "plantilla": PLANTILLA})
    chk("descargar Word con datos corregidos", r.status_code == 200 and "MARIA PRUEBA CORREGIDA" in texto_word(r.content))

    b = ana.get(f"/api/expedientes/{exp_id}/bitacora").json()
    tipos = [e["tipo"] for e in b["eventos"]]
    chk("bitácora: guardar, estados y word", tipos[:1] == ["guardar"] and tipos.count("estado") == 2 and "word" in tipos, tipos)
    guardar_ev = b["eventos"][0]
    chk("bitácora: corrección antes -> después al guardar", guardar_ev["detalle"].get("cambios", {}).get("fecha_escritura") == ["", "10/09/2018"])
    word_ev = [e for e in b["eventos"] if e["tipo"] == "word"][-1]
    chk("bitácora: cambio al editar", word_ev["detalle"].get("cambios", {}).get("nombre_acreditado") == ["MARIA PRUEBA EJEMPLO", "MARIA PRUEBA CORREGIDA"],
        word_ev["detalle"])
    estado_ev = [e for e in b["eventos"] if e["tipo"] == "estado"][0]
    chk("bitácora: estado de -> a", estado_ev["detalle"] == {"de": "GENERADO", "a": "REVISADO"})
    chk("bitácora: versión de plantilla", (b.get("version") or {}).get("huella") == res["datos_extraidos"]["_plantilla_version"])
    chk("beto no ve la bitácora de ana", beto.get(f"/api/expedientes/{exp_id}/bitacora").status_code == 404)

    # ---- Versiones: si la plantilla cambia, el Word del expediente se regenera con la versión original
    ruta = main.resolver_ruta_plantilla(PLANTILLA)
    respaldo_plantilla = ruta + ".respaldo_prueba"
    shutil.copy(ruta, respaldo_plantilla)
    try:
        d = docx.Document(ruta)
        d.add_paragraph("VERSION NUEVA DE LA PLANTILLA")
        d.save(ruta)
        hist = {h["id"]: h for h in ana.get("/api/expedientes").json()}
        chk("Mi Historial avisa que la plantilla cambió", hist[exp_id]["plantilla_cambio"] is True)
        r = ana.post(f"/api/expedientes/{exp_id}/generar-word", json={"datos": datos2, "plantilla": PLANTILLA})
        chk("se regenera con la versión original", r.status_code == 200 and "VERSION NUEVA" not in texto_word(r.content))
        r = ana.post(f"/api/expedientes/{exp_id}/generar-word", json={"datos": datos2, "plantilla": PLANTILLA, "plantilla_actual": True})
        chk("con 'usar la plantilla actual' sale la nueva", r.status_code == 200 and "VERSION NUEVA" in texto_word(r.content))
        hist = {h["id"]: h for h in ana.get("/api/expedientes").json()}
        chk("después ya queda con la versión nueva", hist[exp_id]["plantilla_cambio"] is False)
        # Un expediente nuevo usa la plantilla nueva
        r = ana.post("/api/expedientes/vista-previa", json={"datos": DATOS, "plantilla": PLANTILLA})
        chk("vista previa usa la plantilla actual", "VERSION NUEVA" in texto_word(r.content))
    finally:
        shutil.move(respaldo_plantilla, ruta)

    # ---- Captura manual y coacreditados
    r = ana.post("/api/expedientes/generar-manual", json=dict(DATOS, numero_credito="77" + CREDITO[2:], plantilla=PLANTILLA))
    chk("captura manual", r.status_code == 200 and r.json().get("datos_extraidos", {}).get("_plantilla_version"), r.text[:200])
    coac = dict(DATOS, nombre_acreditado_2="JUAN SEGUNDO EJEMPLO", numero_credito_2="0900203137", numero_carta_2="0524486193",
                fecha_expedicion_2="01/02/2025", fecha_liquidacion_2="03/04/2024", credito_a_salario_2="40.5", monto_credito_2="120000")
    r = ana.post("/api/expedientes/vista-previa", json={"datos": coac, "plantilla": "COAC_EDOMEX_AP"})
    t = texto_word(r.content) if r.status_code == 200 else ""
    chk("coacreditados: Word con los dos trabajadores", "MARIA PRUEBA EJEMPLO" in t and "JUAN SEGUNDO EJEMPLO" in t and "{{" not in t)
    chk("coacreditados: monto del trabajador dos con letra", "$120,000.00 M.N. (CIENTO VEINTE MIL PESOS" in t)
    r = ana.post("/api/expedientes/normalizar-campo", json={"campo": "monto_credito_2", "texto": "5000"}).json()
    chk("normalizar dato del trabajador dos", r["valor"] == "$5,000.00" and "monto_letras_2" in r["extra"])

    # ---- ZIP (usa la versión guardada si el Word ya no está en disco)
    r = ana.post("/api/expedientes/descargar-zip", json=[exp_id])
    chk("descargar ZIP", r.status_code == 200 and len(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == 1)

    # ---- Respaldos
    r = admin.post("/api/admin/respaldos")
    chk("crear respaldo", r.status_code == 200 and r.json().get("id"), r.text[:200])
    lista = admin.get("/api/admin/respaldos").json()["respaldos"]
    chk("lista de respaldos", any(x["id"] == r.json().get("id") for x in lista))
    z = admin.get(f"/api/admin/respaldos/{r.json().get('id')}")
    archivo = zipfile.ZipFile(io.BytesIO(z.content))
    expedientes = json.loads(archivo.read("base/expedientes.json"))
    chk("el respaldo trae los expedientes", any(e["numero_credito"] == CREDITO for e in expedientes))
    chk("el respaldo trae las plantillas", any(n.startswith("plantillas/") for n in archivo.namelist()))
    chk("el respaldo trae las versiones de plantilla", len(json.loads(archivo.read("base/plantilla_versiones.json"))) >= 1)

    # ---- Inactividad / sesión
    r = ana.get("/api/auth/ping", params={"activo": 1}).json()
    chk("ping de sesión", r.get("activa") is True)

# ---- Nombres: orden (nombre primero) y género por el nombre de pila, no por el apellido
import nombres  # noqa: E402
for crudo, carta, esperado, genero in [
    ("PADILLA GUTIERREZ FERNANDO", "", "FERNANDO PADILLA GUTIERREZ", "MASCULINO"),
    ("PADILLA GUTIERREZ FERNANDO", "Fernando Padilla Gutiérrez", "FERNANDO PADILLA GUTIERREZ", "MASCULINO"),
    ("MORALES RUIZ ANDREA", "", "ANDREA MORALES RUIZ", "FEMENINO"),
    ("HERNANDEZ LOPEZ JOSE GUADALUPE", "", "JOSE GUADALUPE HERNANDEZ LOPEZ", "MASCULINO"),
    ("GARCIA CRUZ MARIA DE LOS ANGELES", "", "MARIA DE LOS ANGELES GARCIA CRUZ", "FEMENINO"),
    ("CARLOS ALBERTO SOTO VEGA", "", "CARLOS ALBERTO SOTO VEGA", "MASCULINO"),
    ("MENDOZA PEREZ XOCHITL", "", "XOCHITL MENDOZA PEREZ", "FEMENINO"),
    ("BAUTISTA LUNA ELISA", "", "ELISA BAUTISTA LUNA", "FEMENINO"),
]:
    chk(f"nombre en orden: {crudo}", nombres.ordenar_nombre(crudo, carta) == esperado, nombres.ordenar_nombre(crudo, carta))
    chk(f"género de {crudo}", nombres.genero(crudo) == genero, nombres.genero(crudo))
d = {"nombre_acreditado": "PADILLA GUTIERREZ FERNANDO", "_nombre_carta": "", "_propietarios": [], "datos_inmueble": "ESTADO DE MÉXICO"}
main.aplicar_acto_elegido(d, [])
sug = main.plantilla_sugerida.sugerir_plantilla(d, [], [])
chk("acreditado de la constancia queda con el nombre primero", d["nombre_acreditado"] == "FERNANDO PADILLA GUTIERREZ", d["nombre_acreditado"])
chk("plantilla sugerida de hombre", sug["genero"] == "MASCULINO" and "_H_" in sug["plantilla"], sug["plantilla"])

# ---- Actos de la constancia: escritura "INSTRUMENTO NUMERO", notario con el nombre antes del número, monto en VSMM
import actos_ifrem  # noqa: E402
import services  # noqa: E402
ACTO = ("INSTRUMENTO NUMERO 3404, DE FECHA 31 DE JULIO DE 2001. PASADA ANTE LA FE DEL NOTARIO PUBLICO LIC. VICTOR A VARELA PEREZ "
        "NUMERO 20, DEL DISTRITO DE TOLUCA. EN LA QUE CONSTA EL OTORGAMIENTO DE CREDITO CON CONSTITUCION DE GARANTIA HIPOTECA. "
        "ACREEDOR: INFONAVIT, DEUDOR: NELLY JIMENEZ MARTINEZ. IMPORTE DE LA OBLIGACION GARANTIZADA: 162.9260 VSMM. GARANTIA: INMUEBLE")
da = actos_ifrem.datos_de_acto({"texto": ACTO, "clave": "A.C.S."}, services)
chk("acto: número de escritura (INSTRUMENTO NUMERO)", da.get("numero_escritura_digitos") == "3404", da)
chk("acto: fecha de escritura", da.get("fecha_escritura") == "treinta y uno de julio de dos mil uno", da)
chk("acto: notario con el nombre antes del número", "Victor A. Varela Perez" in str(da.get("notario_origen_completo")) and "veinte" in str(da.get("notario_origen_completo")), da)
chk("acto: monto en veces salario", da.get("credito_a_salario_digitos") == "162.9260", da)
chk("acto: tipo de crédito por la clave A.C.S.", da.get("tipo_credito") == "AP", da)
db = actos_ifrem.datos_de_acto({"texto": "ESCRITURA NUMERO 12,345 DE FECHA 5 DE MARZO DE 2010 OTORGADA ANTE EL NOTARIO PUBLICO NUMERO 96 DEL ESTADO DE MEXICO, LIC. ANA LUISA MARTINEZ SOTO, EN LA QUE CONSTA EL CONTRATO DE MUTUO POR LA CANTIDAD DE $350,000.00"}, services)
chk("acto: notario después del número", "Ana Luisa Martinez Soto" in str(db.get("notario_origen_completo")), db)
chk("acto: monto 'por la cantidad de'", db.get("monto_credito") == "$350,000.00", db)

# ---- Actos: renglón pegado, actos sin volante y plan B por el cierre "VOLANTE ... ANALISTA:"
PEGADO = ["GRAVAMENES, HIPOTECAS, DERECHOS\nREALES Y LIMITACIONES DE DOMINIO\n--- 1000000002505 --- 13/12/2001 A.C.S. INSTRUMENTO NUMERO 3404, "
          "DE FECHA 31 DE JULIO DE 2001. PASADA ANTE LA FE DEL NOTARIO PUBLICO\nLIC. VICTOR A VARELA PEREZ NUMERO 20. VOLANTE: 1000000002505 "
          "CALIFICADOR: ORLANDO BENITEZ RUIZ - ANALISTA:NINGUNO\nVIGENTE"]
ap = actos_ifrem.dividir_actos(PEGADO)
chk("actos: renglón pegado", len(ap) == 1 and ap[0]["volante"] == "1000000002505" and ap[0]["clave"] == "A.C.S.", ap)
SIN_VOLANTE = ["INSCRIPCIONES DE PROPIEDAD\n---\n---\n05/03/1999\nC.V.\nCOMPRAVENTA QUE OTORGA JUAN PEREZ A FAVOR DE NELLY JIMENEZ. "
               "VOLANTE: 1000000001111 CALIFICADOR: X Y - ANALISTA:NINGUNO\n---\n---\n13/12/2001\nA.C.S.\nINSTRUMENTO NUMERO 3404 OTORGAMIENTO "
               "DE CREDITO CON GARANTIA HIPOTECA. VOLANTE: 1000000002505 CALIFICADOR: Z W - ANALISTA:NINGUNO"]
chk("actos: sin volante en la tabla", len(actos_ifrem.dividir_actos(SIN_VOLANTE)) == 2, actos_ifrem.dividir_actos(SIN_VOLANTE))
DESORDEN = ["GRAVAMENES, HIPOTECAS, DERECHOS REALES Y LIMITACIONES DE DOMINIO INSTRUMENTO NUMERO 3404 OTORGAMIENTO DE CREDITO CON GARANTIA "
            "HIPOTECA. VOLANTE: 1000000002505 CALIFICADOR: ORLANDO - ANALISTA:NINGUNO CANCELACION DE HIPOTECA ESCRITURA 5555. VOLANTE: "
            "1000000009999 CALIFICADOR: ANA - ANALISTA:INSCRIPCIONES"]
chk("actos: texto en otro orden (plan B)", len(actos_ifrem.dividir_actos(DESORDEN)) == 2, actos_ifrem.dividir_actos(DESORDEN))
chk("sesión: 15 minutos de inactividad por defecto", main.auth.INACTIVIDAD_SEGUNDOS == int(os.getenv("INACTIVIDAD_SEGUNDOS", "900")))

# ---- Cónyuge: también donde la plantilla dice "con el consentimiento de su cónyuge" sin el nombre
d_cony = main.limpiar_datos_para_plantilla({"nombre_acreditado": "FABIOLA ZETA HERNANDEZ", "numero_credito": CREDITO, "tiene_conyuge": "Sí",
                                            "nombre_conyuge": "ARTURO ESCALONA MACEDO"}, CREDITO)
ruta_cony = os.path.join(main.GENERADOS_DIR, f"prueba_conyuge_{SUF}.docx")
if os.path.exists(main.resolver_ruta_plantilla("EDOMEX_AP_M_CASADA")) and "CASADA" in main.resolver_ruta_plantilla("EDOMEX_AP_M_CASADA"):
    services.generar_word_cancelacion(main.resolver_ruta_plantilla("EDOMEX_AP_M_CASADA"), d_cony, ruta_cony)
    t_cony = texto_word(open(ruta_cony, "rb").read())
    chk("cónyuge en todas las menciones", t_cony.count("ARTURO ESCALONA MACEDO") >= 2 and "cónyuge sobre" not in t_cony, t_cony[:200])
    os.remove(ruta_cony)
import docx as _docx  # noqa: E402
doc_c = _docx.Document()
doc_c.add_paragraph("“EL ACREDITADO”, con el consentimiento de su cónyuge, sobre el inmueble relacionado.")
services.completar_conyuge_en_documento(doc_c, "ANA RUIZ SOTO")
chk("cónyuge: se agrega 'la señora' con su nombre", "cónyuge, la señora ANA RUIZ SOTO, sobre" in doc_c.paragraphs[0].text, doc_c.paragraphs[0].text)

# ---- Editar un expediente guardado y actualizarlo al volver a subirlo
with cliente() as c_ed:
    tok_ed = c_ed.post("/api/auth/login", json={"username": f"ana_{SUF}", "password": "Prueba2026!"}).json().get("token")
    h_ed = {"Authorization": f"Bearer {tok_ed}"}
    cred_ed = "15" + str(uuid.uuid4().int)[:8]
    base_ed = {**DATOS, "numero_credito": cred_ed, "nombre_acreditado": "LUIS PRUEBA EDITAR"}
    r = c_ed.post("/api/expedientes/guardar", json={"expedientes": [{"ref": "e", "datos": base_ed, "plantilla": PLANTILLA}]}, headers=h_ed)
    id_ed = r.json()["resultados"][0]["expediente_id"]
    c_ed.put(f"/api/expedientes/{id_ed}/estado", json={"estado": "ENTREGADO"}, headers=h_ed)
    r = c_ed.put(f"/api/expedientes/{id_ed}/editar", json={"datos": {**base_ed, "nombre_acreditado": "LUIS PRUEBA EDITADO"}, "plantilla": PLANTILLA}, headers=h_ed)
    chk("editar: guarda el cambio", r.status_code == 200 and r.json()["expediente"]["acreditado"] == "LUIS PRUEBA EDITADO", r.text[:300])
    chk("editar: registra qué cambió", "nombre_acreditado" in (r.json().get("cambios") or {}), r.text[:300])
    chk("editar: si estaba Entregado regresa a Generado", r.json()["expediente"]["estado"] == "GENERADO", r.text[:300])
    ev = c_ed.get(f"/api/expedientes/{id_ed}/bitacora", headers=h_ed).json()["eventos"]
    chk("editar: queda en el historial de cambios", any(e["tipo"] == "editar" for e in ev), ev)
    antes = len(c_ed.get("/api/expedientes", headers=h_ed).json())
    r = c_ed.post("/api/expedientes/guardar", json={"expedientes": [{"ref": "u", "actualizar": id_ed, "plantilla": PLANTILLA,
                                                                       "datos": {**base_ed, "nombre_acreditado": "LUIS PRUEBA ACTUALIZADO"}}]}, headers=h_ed)
    res_u = r.json()["resultados"][0]
    chk("volver a subir: actualiza el existente", res_u.get("actualizado") and res_u.get("expediente_id") == id_ed, r.text[:300])
    chk("volver a subir: no crea otro expediente", len(c_ed.get("/api/expedientes", headers=h_ed).json()) == antes)
    ev = c_ed.get(f"/api/expedientes/{id_ed}/bitacora", headers=h_ed).json()["eventos"]
    chk("volver a subir: queda en el historial de cambios", any(e["tipo"] == "actualizar" for e in ev), ev)
    fila = next(h for h in c_ed.get("/api/expedientes", headers=h_ed).json() if h["id"] == id_ed)
    chk("historial: trae la plantilla con la que se guardó", fila["datos_extraidos"].get("plantilla_seleccionada") == PLANTILLA, fila["datos_extraidos"].get("plantilla_seleccionada"))
    tok_b = c_ed.post("/api/auth/login", json={"username": f"beto_{SUF}", "password": "Prueba2026!"}).json().get("token")
    r = c_ed.put(f"/api/expedientes/{id_ed}/editar", json={"datos": base_ed}, headers={"Authorization": f"Bearer {tok_b}"})
    chk("editar: otro operador no puede editarlo", r.status_code == 404, r.status_code)

# ---- Contraseñas: el admin cambia la de un usuario escribiendo la contraseña actual de ese usuario
with cliente() as c_pw:
    h_adm = {"Authorization": "Bearer " + c_pw.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    beto_id = next(u["id"] for u in c_pw.get("/api/admin/usuarios", headers=h_adm).json() if u["username"] == f"beto_{SUF}")
    r = c_pw.post(f"/api/admin/usuarios/{beto_id}/password", json={"password_nueva": "Nueva2026!"}, headers=h_adm)
    chk("admin: sin la contraseña actual del usuario no la cambia", r.status_code == 400, r.status_code)
    r = c_pw.post(f"/api/admin/usuarios/{beto_id}/password", json={"password_actual": os.getenv("ADMIN_PASSWORD", "admin123"), "password_nueva": "Nueva2026!"}, headers=h_adm)
    chk("admin: con su propia contraseña (no la del usuario) no la cambia", r.status_code == 400, r.status_code)
    r = c_pw.post(f"/api/admin/usuarios/{beto_id}/password", json={"password_actual": "Prueba2026!", "password_nueva": "Nueva2026!"}, headers=h_adm)
    chk("admin: con la contraseña actual del usuario sí la cambia", r.status_code == 200, r.text[:200])
    chk("el usuario entra con la nueva", c_pw.post("/api/auth/login", json={"username": f"beto_{SUF}", "password": "Nueva2026!"}).status_code == 200)
    import auth as _auth  # noqa: E402

    class _U:  # usuario de prueba para el saludo
        def __init__(self, n, u): self.nombre_mostrar, self.username = n, u
    chk("saludo: Bienvenida para Laura Méndez", _auth.genero_de_usuario(_U("Laura Méndez", "laura")) == "FEMENINO")
    chk("saludo: Bienvenido para Pedro Salinas", _auth.genero_de_usuario(_U("Pedro Salinas", "pedro")) == "MASCULINO")

# ---- Plantilla guardada: descargar desde Mi Historial no debe cambiarla, y la corrección la regresa
with cliente() as c_pl:
    import reparaciones as _rep  # noqa: E402
    from database import SessionLocal as _SL  # noqa: E402
    h_b = {"Authorization": "Bearer " + c_pl.post("/api/auth/login", json={"username": f"beto_{SUF}", "password": "Nueva2026!"}).json()["token"]}
    cred = "19" + uuid.uuid4().hex[:8].translate(str.maketrans("abcdef", "123456"))
    r = c_pl.post("/api/expedientes/guardar", headers=h_b, json={"expedientes": [{"ref": "p1", "plantilla": "EDOMEX_AP_M_SOLTERA",
        "datos": {"numero_credito": cred, "nombre_acreditado": "MARIA PRUEBA PLANTILLA"}}]})
    eid = r.json()["resultados"][0]["expediente_id"]
    r = c_pl.post(f"/api/expedientes/{eid}/generar-word", headers=h_b, json={"datos": {"numero_credito": cred}})
    _db = _SL()
    from models import Expediente as _E  # noqa: E402
    _p = lambda: (_db.get(_E, uuid.UUID(eid)).datos_extraidos or {}).get("plantilla_seleccionada")  # noqa: E731
    chk("descarga sin plantilla usa la que se guardó", r.status_code == 200 and _p() == "EDOMEX_AP_M_SOLTERA", _p())
    # Así descargaba la pantalla anterior desde Mi Historial (con la plantilla por defecto)
    c_pl.post(f"/api/expedientes/{eid}/generar-word", headers=h_b, json={"plantilla": "CDMX_AP_H_SOLTERO", "datos": {"numero_credito": cred}})
    _db.expire_all()
    chk("(simulación) quedó con la plantilla por defecto", _p() == "CDMX_AP_H_SOLTERO", _p())
    res = _rep.reparar_plantillas(_db)
    _db.expire_all()
    chk("corrección: regresa a la plantilla con la que se guardó", _p() == "EDOMEX_AP_M_SOLTERA", _p())
    chk("corrección: lo reporta", any(x["id"] == eid for x in res["detalle"]), res["corregidos"])
    res2 = _rep.reparar_plantillas(_db, simular=True)
    chk("corrección: la segunda vez ya no cambia nada de ese expediente", not any(x["id"] == eid for x in res2["detalle"]))
    _db.close()

# ---- Eliminar un usuario que tiene expedientes (con respaldo previo)
with cliente() as c_del:
    h_adm = {"Authorization": "Bearer " + c_del.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    nombre_del = f"borrar_{SUF}"
    c_del.post("/api/admin/usuarios", data={"username": nombre_del, "password": "Prueba2026!", "es_admin": "false"}, headers=h_adm)
    h_del = {"Authorization": "Bearer " + c_del.post("/api/auth/login", json={"username": nombre_del, "password": "Prueba2026!"}).json()["token"]}
    for i in range(2):
        c_del.post("/api/expedientes/guardar", headers=h_del, json={"expedientes": [{"ref": f"d{i}", "plantilla": "CDMX_AP_H_SOLTERO",
            "datos": {"numero_credito": f"18{i}{SUF[:6]}".ljust(10, "0")[:10], "nombre_acreditado": "JUAN BORRAR"}}]})
    # (el login del usuario deja su cookie: se vuelve a entrar como admin)
    h_adm = {"Authorization": "Bearer " + c_del.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    lista = c_del.get("/api/admin/usuarios", headers=h_adm).json()
    u_del = next(u for u in lista if u["username"] == nombre_del)
    chk("usuarios: dice cuántos expedientes tiene", u_del.get("expedientes") == 2, u_del.get("expedientes"))
    r = c_del.delete(f"/api/admin/usuarios/{u_del['id']}", headers=h_adm)
    chk("eliminar: con expedientes y sin confirmar, no lo borra (409)", r.status_code == 409, r.status_code)
    r = c_del.delete(f"/api/admin/usuarios/{u_del['id']}", params={"con_expedientes": "true", "confirmar": "otro"}, headers=h_adm)
    chk("eliminar: confirmando con otro nombre, no lo borra", r.status_code == 400, r.status_code)
    n_resp = len(c_del.get("/api/admin/respaldos", headers=h_adm).json())
    r = c_del.delete(f"/api/admin/usuarios/{u_del['id']}", params={"con_expedientes": "true", "confirmar": nombre_del}, headers=h_adm)
    chk("eliminar: con sus expedientes", r.status_code == 200 and r.json().get("expedientes_eliminados") == 2, r.text[:200])
    chk("eliminar: hace un respaldo antes", bool(r.json().get("respaldo_id")))
    from database import SessionLocal as _SL2  # noqa: E402
    import models as _m  # noqa: E402
    _d = _SL2()
    chk("eliminar: ya no quedan sus expedientes", _d.query(_m.Expediente).filter(_m.Expediente.usuario_propietario == nombre_del).count() == 0)
    chk("eliminar: ya no queda su actividad", _d.query(_m.Actividad).filter(_m.Actividad.username == nombre_del).count() == 0)
    _d.close()
    chk("eliminar: ya no puede entrar", c_del.post("/api/auth/login", json={"username": nombre_del, "password": "Prueba2026!"}).status_code in (400, 401, 403))

# ---- Quitar permisos de administrador a otro admin (para poder eliminarlo)
with cliente() as c_rol:
    h_adm = {"Authorization": "Bearer " + c_rol.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    nombre_adm2 = f"jefa_{SUF}"
    c_rol.post("/api/admin/usuarios", data={"username": nombre_adm2, "password": "Prueba2026!", "es_admin": "true"}, headers=h_adm)
    u2 = next(u for u in c_rol.get("/api/admin/usuarios", headers=h_adm).json() if u["username"] == nombre_adm2)
    chk("rol: se creó como administrador", u2["es_admin"] is True)
    r = c_rol.delete(f"/api/admin/usuarios/{u2['id']}", headers=h_adm)
    chk("rol: un administrador no se puede eliminar", r.status_code == 400, r.status_code)
    r = c_rol.put(f"/api/admin/usuarios/{u2['id']}/rol", json={"es_admin": False}, headers=h_adm)
    chk("rol: se le quitan los permisos de administrador", r.status_code == 200 and r.json()["es_admin"] is False, r.text[:200])
    adm = next(u for u in c_rol.get("/api/admin/usuarios", headers=h_adm).json() if u["username"] == "admin")
    r = c_rol.put(f"/api/admin/usuarios/{adm['id']}/rol", json={"es_admin": False}, headers=h_adm)
    chk("rol: la cuenta principal admin no se puede cambiar", r.status_code == 400, r.status_code)
    r = c_rol.delete(f"/api/admin/usuarios/{u2['id']}", headers=h_adm)
    chk("rol: ya como operador sí se puede eliminar", r.status_code == 200, r.text[:200])

# ---- Editar desde Mi Historial con el espacio de trabajo (guardar con actualizar + tipo_edicion)
with cliente() as c_ed:
    h_b = {"Authorization": "Bearer " + c_ed.post("/api/auth/login", json={"username": f"beto_{SUF}", "password": "Nueva2026!"}).json()["token"]}
    cred = "17" + SUF[:6].translate(str.maketrans("abcdef", "123456")).ljust(8, "0")
    r = c_ed.post("/api/expedientes/guardar", headers=h_b, json={"expedientes": [{"ref": "e1", "plantilla": "CDMX_AP_M_SOLTERA",
        "datos": {"numero_credito": cred, "nombre_acreditado": "ROSA EDICION", "folio_real": "111"}}]})
    eid = r.json()["resultados"][0]["expediente_id"]
    r = c_ed.post("/api/expedientes/guardar", headers=h_b, json={"expedientes": [{"ref": "e1", "plantilla": "EDOMEX_AP_M_SOLTERA", "actualizar": eid,
        "tipo_edicion": "editar", "datos": {"numero_credito": cred, "nombre_acreditado": "ROSA EDICION", "folio_real": "222"}}]})
    res = r.json()
    chk("edición completa: actualiza el mismo expediente", res.get("actualizados") == 1 and res["resultados"][0]["expediente_id"] == eid, str(res)[:200])
    b = c_ed.get(f"/api/expedientes/{eid}/bitacora", headers=h_b).json()
    tipos = [e.get("tipo") for e in b.get("eventos", [])]
    chk("edición completa: queda como «editó» en su historial", "editar" in tipos, tipos)
    chk("edición completa: guarda la plantilla nueva", b.get("plantilla") == "EDOMEX_AP_M_SOLTERA", b.get("plantilla"))

# ---- Compartir expedientes con usuarios de revisión
with cliente() as c_sh:
    def entrar_como(u, p):
        tok = c_sh.post("/api/auth/login", json={"username": u, "password": p}).json()["token"]
        c_sh.cookies.clear()  # cada petición va con su propio token (sin la cookie del último login)
        return {"Authorization": "Bearer " + tok}
    h_adm = entrar_como("admin", os.getenv("ADMIN_PASSWORD", "admin123"))
    prov, rev, otro = f"proy_{SUF}", f"rev_{SUF}", f"otro_{SUF}"
    c_sh.post("/api/admin/usuarios", data={"username": prov, "password": "Prueba2026!", "rol": "operador"}, headers=h_adm)
    c_sh.post("/api/admin/usuarios", data={"username": rev, "password": "Prueba2026!", "rol": "revisor"}, headers=h_adm)
    c_sh.post("/api/admin/usuarios", data={"username": otro, "password": "Prueba2026!"}, headers=h_adm)
    lista_u = {u["username"]: u for u in c_sh.get("/api/admin/usuarios", headers=h_adm).json()}
    chk("revisor: se crea con su tipo", lista_u[rev]["rol"] == "revisor" and lista_u[prov]["rol"] == "operador", lista_u[rev].get("rol"))
    h_p = entrar_como(prov, "Prueba2026!")
    ids = []
    for i in range(3):
        r = c_sh.post("/api/expedientes/guardar", headers=h_p, json={"expedientes": [{"ref": f"s{i}", "plantilla": "CDMX_AP_H_SOLTERO",
            "datos": {"numero_credito": f"16{i}{SUF[:7]}".ljust(10, "1")[:10], "nombre_acreditado": f"COMPARTIDO {i}"}}]})
        ids.append(r.json()["resultados"][0]["expediente_id"])
    h_r = entrar_como(rev, "Prueba2026!")
    chk("revisor: entra con rol revisor", c_sh.get("/api/auth/me", headers=h_r).json()["user"]["role"] == "revisor")
    chk("revisor: sin compartir no ve nada", c_sh.get("/api/compartidos", headers=h_r).json() == [])
    chk("revisor: no ve expedientes ajenos", c_sh.get(f"/api/expedientes/{ids[0]}/bitacora", headers=h_r).status_code == 404)
    r = c_sh.post("/api/expedientes/guardar", headers=h_r, json={"expedientes": [{"ref": "x", "plantilla": "CDMX_AP_H_SOLTERO", "datos": {"numero_credito": "1234567890"}}]})
    chk("revisor: no puede crear expedientes", r.status_code == 403, r.status_code)
    r = c_sh.post("/api/expedientes/generar-manual", headers=h_r, json={"datos": {"numero_credito": "1234567890"}, "plantilla": "CDMX_AP_H_SOLTERO"})
    chk("revisor: no puede usar captura manual", r.status_code == 403, r.status_code)
    chk("revisor: no puede compartir", c_sh.get("/api/compartir/directorio", headers=h_r).status_code == 403)
    dire = c_sh.get("/api/compartir/directorio", headers=h_p).json()
    chk("directorio: operadores, revisores y administradores", any(d["usuario"] == otro for d in dire) and any(d["usuario"] == "admin" and d["rol"] == "admin" for d in dire), str(dire)[:160])
    r = c_sh.post("/api/compartir", headers=h_p, json={"ids": ids[:1], "usuarios": ["admin"]})
    chk("compartir: también con un administrador", r.status_code == 200 and r.json()["nuevos"] == 1, r.text[:200])
    chk("compartir: el administrador lo ve en «Compartidos conmigo»", {x["id"] for x in c_sh.get("/api/compartidos", headers=h_adm).json()} >= {ids[0]})
    c_sh.post("/api/compartir/quitar", headers=h_p, json={"ids": ids[:1], "usuarios": ["admin"]})
    chk("compartir: se le quita el acceso al administrador", ids[0] not in {x["id"] for x in c_sh.get("/api/compartidos", headers=h_adm).json()})
    # Compartir dos expedientes
    r = c_sh.post("/api/compartir", headers=h_p, json={"ids": ids[:2], "usuarios": [rev]})
    chk("compartir: dos expedientes", r.status_code == 200 and r.json()["nuevos"] == 2, r.text[:200])
    vistos = {x["id"] for x in c_sh.get("/api/compartidos", headers=h_r).json()}
    chk("compartidos: el revisor ve esos dos", vistos == set(ids[:2]), vistos)
    chk("compartidos: no ve el tercero", c_sh.get(f"/api/expedientes/{ids[2]}/bitacora", headers=h_r).status_code == 404)
    r = c_sh.put(f"/api/expedientes/{ids[0]}/estado", headers=h_r, json={"estado": "REVISADO"})
    chk("revisor: cambia el estado", r.status_code == 200 and r.json()["estado"] == "REVISADO", r.text[:200])
    r = c_sh.post(f"/api/expedientes/{ids[0]}/generar-word", headers=h_r, json={"datos": {}})
    chk("revisor: descarga el Word", r.status_code == 200 and len(r.content) > 1000, r.status_code)
    r = c_sh.post("/api/expedientes/guardar", headers=h_r, json={"expedientes": [{"ref": "e", "actualizar": ids[0], "tipo_edicion": "editar",
        "plantilla": "CDMX_AP_H_SOLTERO", "datos": {"numero_credito": f"160{SUF[:7]}".ljust(10, "1")[:10], "nombre_acreditado": "COMPARTIDO EDITADO"}}]})
    chk("revisor: edita un compartido", r.status_code == 200 and r.json().get("actualizados") == 1, r.text[:200])
    from database import SessionLocal as _SL3  # noqa: E402
    import models as _m3  # noqa: E402
    _d3 = _SL3()
    chk("compartido: sigue siendo de quien lo hizo", _d3.get(_m3.Expediente, uuid.UUID(ids[0])).usuario_propietario == prov)
    _d3.close()
    hist_p = {x["id"]: x for x in c_sh.get("/api/expedientes", headers=h_p).json()}
    chk("Mi Historial: dice con quién está compartido", [g["usuario"] for g in hist_p[ids[0]]["compartido_con"]] == [rev], hist_p[ids[0]].get("compartido_con"))
    acc = c_sh.get(f"/api/compartir/expediente/{ids[0]}", headers=h_p).json()
    chk("accesos de un expediente", len(acc) == 1 and acc[0]["usuario"] == rev and acc[0]["directo"], acc)
    tipos = [e["tipo"] for e in c_sh.get(f"/api/expedientes/{ids[0]}/bitacora", headers=h_p).json()["eventos"]]
    chk("bitácora: compartir, estado y edición del revisor", "compartir" in tipos and "estado" in tipos and "editar" in tipos, tipos)
    # Otro operador no puede compartir lo ajeno
    h_o = entrar_como(otro, "Prueba2026!")
    r = c_sh.post("/api/compartir", headers=h_o, json={"ids": ids, "usuarios": [rev]})
    chk("compartir: nadie comparte expedientes ajenos", r.status_code == 400, r.status_code)
    # Quitar acceso
    c_sh.post("/api/compartir/quitar", headers=h_p, json={"ids": [ids[1]], "usuarios": [rev]})
    vistos = {x["id"] for x in c_sh.get("/api/compartidos", headers=h_r).json()}
    chk("quitar acceso: ya no lo ve", vistos == {ids[0]}, vistos)
    # Compartir todo (incluye los futuros): «otro» se vuelve revisor
    c_sh.put(f"/api/admin/usuarios/{lista_u[otro]['id']}/rol", headers=h_adm, json={"rol": "revisor"})
    h_o = entrar_como(otro, "Prueba2026!")
    c_sh.post("/api/compartir", headers=h_p, json={"todos": True, "usuarios": [otro]})
    r = c_sh.post("/api/expedientes/guardar", headers=h_p, json={"expedientes": [{"ref": "n", "plantilla": "CDMX_AP_H_SOLTERO",
        "datos": {"numero_credito": f"169{SUF[:7]}".ljust(10, "1")[:10], "nombre_acreditado": "NUEVO DESPUES"}}]})
    nuevo_id = r.json()["resultados"][0]["expediente_id"]
    vistos_o = {x["id"] for x in c_sh.get("/api/compartidos", headers=h_o).json()}
    chk("compartir todo: ve todos, también los nuevos", set(ids) | {nuevo_id} <= vistos_o, len(vistos_o))
    chk("compartir todo: aparece en la lista", [x["usuario"] for x in c_sh.get("/api/compartir/todos", headers=h_p).json()] == [otro])
    c_sh.post("/api/compartir/quitar", headers=h_p, json={"todos": True, "usuarios": [otro]})
    chk("quitar todo: ya no ve nada", c_sh.get("/api/compartidos", headers=h_o).json() == [])
    # Cambiar el tipo de usuario
    r = c_sh.put(f"/api/admin/usuarios/{lista_u[rev]['id']}/rol", headers=h_adm, json={"rol": "operador"})
    chk("tipo de usuario: revisor -> operador", r.status_code == 200 and r.json()["rol"] == "operador", r.text[:150])

# ---- Restablecer una contraseña olvidada (el admin confirma con la suya; el usuario debe cambiarla al entrar)
with cliente() as c_rp:
    def entrar_rp(u, p):
        r = c_rp.post("/api/auth/login", json={"username": u, "password": p})
        c_rp.cookies.clear()
        return r
    h_adm = {"Authorization": "Bearer " + entrar_rp("admin", os.getenv("ADMIN_PASSWORD", "admin123")).json()["token"]}
    olv = f"olvido_{SUF}"
    c_rp.post("/api/admin/usuarios", data={"username": olv, "password": "Original2026!"}, headers=h_adm)
    uid = next(u["id"] for u in c_rp.get("/api/admin/usuarios", headers=h_adm).json() if u["username"] == olv)
    r = c_rp.post(f"/api/admin/usuarios/{uid}/restablecer-password", headers=h_adm, json={"password_admin": "mala", "password_nueva": "Temporal2026!"})
    chk("restablecer: con contraseña de admin incorrecta no", r.status_code == 400, r.status_code)
    r = c_rp.post(f"/api/admin/usuarios/{uid}/restablecer-password", headers=h_adm, json={"password_admin": os.getenv("ADMIN_PASSWORD", "admin123"), "password_nueva": "Temporal2026!"})
    chk("restablecer: sin saber la contraseña del usuario", r.status_code == 200, r.text[:200])
    chk("restablecer: la anterior ya no sirve", entrar_rp(olv, "Original2026!").status_code == 401)
    r = entrar_rp(olv, "Temporal2026!")
    chk("restablecer: entra con la temporal y se le pide cambiarla", r.status_code == 200 and r.json()["user"]["debe_cambiar_password"] is True, r.text[:200])
    h_o = {"Authorization": "Bearer " + r.json()["token"]}
    r = c_rp.post("/api/auth/cambiar-password", headers=h_o, json={"password_actual": "Temporal2026!", "password_nueva": "Propia2026!"})
    chk("restablecer: la cambia por una propia", r.status_code == 200, r.text[:200])
    r = entrar_rp(olv, "Propia2026!")
    chk("restablecer: ya no se le pide cambiarla", r.status_code == 200 and r.json()["user"]["debe_cambiar_password"] is False)

# --- Datos del inmueble sin «:» ---
_arm = main._texto_inmueble({"1_vivienda": "CASA 14", "4_lote": "14", "5_manzana": "15", "13_colonia": "LOS ALAMOS"})
chk("inmueble: se arma sin dos puntos", ":" not in _arm and "Manzana 15" in _arm, _arm)
_l1 = main.limpiar_datos_para_plantilla({"datos_inmueble": _arm, "numero_credito": "1"})["datos_inmueble"]
chk("inmueble: al volver a limpiarlo no se pierde nada", _l1 == _arm, _l1)
_l2 = main.limpiar_datos_para_plantilla({"datos_inmueble": "Vivienda: CASA 14, Manzana: 15", "numero_credito": "1"})["datos_inmueble"]
chk("inmueble: guardado antes con dos puntos sale sin ellos", _l2 == "Vivienda CASA 14, Manzana 15", _l2)

# --- Montos escritos de muchas formas ---
for _txt, _esp in [("$1,250,300.00", 1250300.0), ("$1´250,300.00", 1250300.0), ("$1'250,300.00", 1250300.0), ("$1.800.000.00", 1800000.0),
                   ("$1 250 300.00", 1250300.0), ("$170,869 02 M.N.", 170869.02), ("$250000", 250000.0), ("$198,640.00", 198640.0)]:
    chk(f"monto: {_txt}", services.normalizar_monto(_txt) == _esp, services.normalizar_monto(_txt))
chk("monto con letra de $1.800.000.00", services.numero_a_letras("$1.800.000.00").startswith("UN MILLÓN OCHOCIENTOS MIL PESOS"))
_d = actos_ifrem.datos_de_acto({"texto": "IMPORTE DE LA OBLIGACION GARANTIZADA: $170,869 02 M.N., EQUIVALENTE A 106.8765 VSMM.", "clave": "A.C.S."}, services)
chk("monto del acto con centavos separados por espacio", _d.get("monto_credito") == "$170,869.02", _d.get("monto_credito"))

# --- Número de escritura en distintas redacciones ---
for _txt in ["TESTIMONIO DE LA ESCRITURA PUBLICA NUMERO 6511, DE FECHA", "ESCRITURA No. 6,511 DE FECHA", "ESC. 6511 DEL",
             "ESCRITURA PUBLICA PASADA NUMERO 6511, ANTE", "INSTRUMENTO NOTARIAL 6511", "NUMERO DE ESCRITURA: 6511",
             "ESCRITURA NUMERO SEIS MIL QUINIENTOS ONCE, DE FECHA", "ESCRlTURA PUBLICA NUMERO 65I1, DE FECHA", "ESCRITURA PUBLICA N° 6511"]:
    chk(f"escritura: {_txt[:40]}", actos_ifrem.numero_escritura_de_texto(_txt) == "6511", actos_ifrem.numero_escritura_de_texto(_txt))
chk("escritura: no confunde el número del notario", actos_ifrem.numero_escritura_de_texto("TESTIMONIO NOTARIAL, NOTARIO PUBLICO NUMERO 68") is None)

# --- Notario: las redacciones más comunes de las constancias ---
_ESP_NOT = "del Licenciado Carlos Lara Vazquez, Notario Público número sesenta y siete del Estado de México"
for _txt in [
    "PASADA ANTE LA FE DEL NOTARIO PÚBLICO NUMERO 67 LIC. CARLOS LARA VAZQUEZ CON RESIDENCIA EN TOLUCA ESTADO DE MEXICO. EN EL QUE",
    "PASADA ANTE LA FE DEL NOTARIO PÚBLICO NUMERO 67 LIC.CARLOS LARA VAZQUEZ CON RESIDENCIA EN TOLUCA ESTADO DE MEXICO.",
    "ANTE LA FE DEL NOTARIO PUBLICO No. 67 LICENCIADO CARLOS LARA VAZQUEZ, CON RESIDENCIA EN TOLUCA, ESTADO DE MEXICO,",
    "ANTE LA FE DEL LIC. CARLOS LARA VAZQUEZ, NOTARIO PUBLICO NUMERO 67 DEL ESTADO DE MEXICO,",
    "OTORGADA ANTE EL NOTARIO PUBLICO NUMERO 67 DEL ESTADO DE MEXICO, LIC. CARLOS LARA VAZQUEZ, EN LA QUE",
    "ANTE LA NOTARIA PUBLICA NUMERO 67 A CARGO DEL LIC. CARLOS LARA VAZQUEZ, DEL ESTADO DE MEXICO",
    "ANTE LA FE DEL NOTARIO PUBLICO NUMERO SESENTA Y SIETE LIC. CARLOS LARA VAZQUEZ CON RESIDENCIA EN TOLUCA",
    "ANTE LA FE DEL NOT. PUB. NUM. 67 LIC. CARLOS LARA VAZQUEZ, CON RESIDENCIA EN TOLUCA",
    "ANTE LA FE DEL NOTARIO PUBLICO NUMERO 67 DE TOLUCA, ESTADO DE MEXICO, LIC. CARLOS LARA VAZQUEZ.",
]:
    chk(f"notario: {_txt[:45]}", actos_ifrem.notario_de_acto(_txt, services) == _ESP_NOT, actos_ifrem.notario_de_acto(_txt, services))
_n = actos_ifrem.notario_de_acto("ANTE LA FE DEL NOTARIO PUBLICO NUMERO 120 LIC. JORGE PEREZ GOMEZ CON RESIDENCIA EN LA CIUDAD DE MEXICO.", services) or ""
chk("notario: de la Ciudad de México por su residencia", _n.endswith("de la Ciudad de México"), _n)
_n = actos_ifrem.notario_de_acto("ANTE LA FE DE LA LICENCIADA ANA LUISA SOTO RUIZ, NOTARIA PUBLICA NUMERO 5 DEL ESTADO DE MEXICO,", services) or ""
chk("notario: licenciada", _n == "de la Licenciada Ana Luisa Soto Ruiz, Notario Público número cinco del Estado de México", _n)
chk("notario: guardado antes como «Notaria Publica» se corrige", services.acentos_notario("de la Licenciada Ana Soto, Notaria Publica numero cinco del Estado de Mexico") == "de la Licenciada Ana Soto, Notario Público número cinco del Estado de México")
chk("personas: sin «LA SEÑORA»", actos_ifrem.separar_personas("ADRIAN MANUEL REBOLLO PEDRAZA Y LA SEÑORA TERESA MORALES CORONA") == ["ADRIAN MANUEL REBOLLO PEDRAZA", "TERESA MORALES CORONA"])
chk("personas: sin «SU ESPOSA LA C.»", actos_ifrem.separar_personas("JUAN PEREZ LOPEZ Y SU ESPOSA LA C. MARIA RUIZ SOTO") == ["JUAN PEREZ LOPEZ", "MARIA RUIZ SOTO"])

# --- Foto de portada del perfil ---
with cliente() as c_por:
    _h = {"Authorization": "Bearer " + c_por.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    c_por.cookies.clear()
    _img = "data:image/jpeg;base64," + "QUJD" * 10
    _r = c_por.put("/api/perfil/portada", headers=_h, json={"portada": _img, "posicion": 30})
    chk("portada: se guarda", _r.status_code == 200 and _r.json()["portada"] == _img and _r.json()["portada_pos"] == 30, _r.text[:200])
    _r = c_por.put("/api/perfil/portada", headers=_h, json={"posicion": 140})
    chk("portada: acomodar (se limita a 100)", _r.status_code == 200 and _r.json()["portada_pos"] == 100 and _r.json()["portada"] == _img, _r.text[:200])
    chk("portada: no imagen se rechaza", c_por.put("/api/perfil/portada", headers=_h, json={"portada": "data:text/html;base64,QUJD"}).status_code == 400)
    chk("portada: muy grande se rechaza", c_por.put("/api/perfil/portada", headers=_h, json={"portada": "data:image/jpeg;base64," + "A" * 1_700_000}).status_code == 400)
    chk("portada: sin sesión no", c_por.put("/api/perfil/portada", json={"posicion": 10}).status_code in (401, 403))
    _r = c_por.delete("/api/perfil/portada", headers=_h)
    chk("portada: quitar", _r.status_code == 200 and _r.json()["portada"] == "" and _r.json()["portada_pos"] == 50, _r.text[:200])

# --- Al editar, lo escrito a mano se imprime tal cual y las conversiones siguen al número ---
_L = main.limpiar_datos_para_plantilla
_INM = 'la vivienda habitacional marcada con el número dieciocho guion "a", construido sobre el lote dieciocho, manzana once, del lote condominal tres'
_RAW = "DATOS DE IDENTIFICACIÓN: VIVIENDA: CASA 18-A LOTE: 18 MANZANA: 11 MUNICIPIO: TOLUCA SUPERFICIE: 60 M2"
chk("inmueble escrito a mano: se respeta", _L({"datos_inmueble": _INM, "texto_raw": _RAW})["datos_inmueble"] == _INM)
chk("inmueble escrito a mano: no se recorta a lote y manzana", _L({"datos_inmueble": _INM})["datos_inmueble"] == _INM)
chk("inmueble leído del PDF: se arma con la constancia", _L({"datos_inmueble": "x", "texto_raw": _RAW}, desde_pdf=True)["datos_inmueble"].startswith("Vivienda CASA 18-A, Lote 18"))
chk("folio: al cambiar el número cambian sus letras", _L({"folio_real": '"321250" (TRESCIENTOS VEINTIÚN MIL DOSCIENTOS CUARENTA)'})["folio_real"] == '"321250" (TRESCIENTOS VEINTIÚN MIL DOSCIENTOS CINCUENTA)')
chk("carta: al cambiar el número cambian sus letras", _L({"numero_carta": '"0543910958" (cero cinco cuatro tres nueve uno cero nueve cinco nueve)'})["numero_carta"].endswith("cinco ocho)"))
chk("carta 2: igual", _L({"numero_carta_2": "0543910958 (cero cinco)", "nombre_acreditado_2": "X"})["numero_carta_2"] == '"0543910958" (cero cinco cuatro tres nueve uno cero nueve cinco ocho)')
chk("escritura: número con sus letras viejas", _L({"numero_escritura": "15813 (quince mil ochocientos doce)"})["numero_escritura"] == "quince mil ochocientos trece")
chk("VSM con número: a letra", _L({"credito_a_salario": "78.6457"})["credito_a_salario"] == "setenta y ocho punto seis mil cuatrocientos cincuenta y siete")
chk("VSM con letra: tal cual", _L({"credito_a_salario": "setenta y ocho punto cinco"})["credito_a_salario"] == "setenta y ocho punto cinco")
chk("monto con letra corregido a mano (misma cantidad): se respeta",
    _L({"monto_credito": "$125,733.76", "monto_letras": "CIENTO VEINTICINCO MIL SETECIENTOS TREINTA Y TRES PESOS 76/100 M.N."})["monto_letras"].endswith("76/100 M.N."))
chk("monto con letra de otro monto: se recalcula",
    _L({"monto_credito": "$125,733.76", "monto_letras": "CIEN PESOS CON CERO CENTAVOS, MONEDA NACIONAL"})["monto_letras"].startswith("CIENTO VEINTICINCO MIL"))
chk("monto con letra: de los PDF siempre se calcula",
    _L({"monto_credito": "$125,733.76", "monto_letras": "CIENTO VEINTICINCO MIL SETECIENTOS TREINTA Y TRES PESOS 76/100 M.N."}, desde_pdf=True)["monto_letras"].endswith("MONEDA NACIONAL"))
chk("desglose del inmueble: se conserva", _L({"datos_inmueble": _INM, "4_lote": "18"}).get("4_lote") == "18")
with cliente() as c_cv:
    _h = {"Authorization": "Bearer " + c_cv.post("/api/auth/login", json={"username": f"ana_{SUF}", "password": "Prueba2026!"}).json()["token"]}
    c_cv.cookies.clear()
    _cred = "16" + str(uuid.uuid4().int)[:8]
    _base = {**DATOS, "numero_credito": _cred}
    _id = c_cv.post("/api/expedientes/guardar", json={"expedientes": [{"ref": "c", "datos": _base, "plantilla": PLANTILLA}]}, headers=_h).json()["resultados"][0]["expediente_id"]
    _ed = {**_base, "datos_inmueble": _INM, "folio_real": '"321250" (TRESCIENTOS VEINTIÚN MIL DOSCIENTOS CUARENTA)'}
    r = c_cv.put(f"/api/expedientes/{_id}/editar", json={"datos": _ed, "plantilla": PLANTILLA}, headers=_h)
    chk("editar inmueble a mano: se guarda", r.status_code == 200, r.text[:200])
    r = c_cv.post(f"/api/expedientes/{_id}/generar-word", json={"datos": {}, "plantilla": PLANTILLA}, headers=_h)
    _t = texto_word(r.content) if r.status_code == 200 else ""
    chk("Word: el inmueble escrito a mano", "marcada con el número dieciocho guion" in _t and "del lote condominal tres" in _t, _t[:300])
    chk("Word: el folio con sus letras nuevas", "321250" in _t and "DOSCIENTOS CINCUENTA" in _t.upper(), _t[:300])
    r = c_cv.post("/api/expedientes/normalizar-campo", headers=_h, json={"campo": "folio_real", "texto": "00321240"}).json()
    chk("normalizar folio: con su formato", r["valor"] == '"321240" (TRESCIENTOS VEINTIÚN MIL DOSCIENTOS CUARENTA)', r)

# --- El administrador ve el perfil de otro usuario y le cierra sesiones ---
with cliente() as c_vp:
    _ha = {"Authorization": "Bearer " + c_vp.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    c_vp.cookies.clear()
    _tok_ana = c_vp.post("/api/auth/login", json={"username": f"ana_{SUF}", "password": "Prueba2026!"}).json()["token"]
    c_vp.cookies.clear()
    _ho = {"Authorization": "Bearer " + _tok_ana}
    _uid = next(u["id"] for u in c_vp.get("/api/admin/usuarios", headers=_ha).json() if u["username"] == f"ana_{SUF}")
    r = c_vp.get(f"/api/admin/usuarios/{_uid}/perfil", headers=_ha)
    _p = r.json() if r.status_code == 200 else {}
    chk("perfil ajeno: el admin lo ve", r.status_code == 200 and _p.get("username") == f"ana_{SUF}" and "resumen" in _p and _p.get("rol") == "operador", r.text[:200])
    chk("perfil ajeno: con sus sesiones", len(_p.get("sesiones") or []) >= 1, _p.get("sesiones"))
    chk("perfil ajeno: un operador no puede", c_vp.get(f"/api/admin/usuarios/{_uid}/perfil", headers=_ho).status_code == 403)
    chk("perfil ajeno: id inválido", c_vp.get("/api/admin/usuarios/no-existe/perfil", headers=_ha).status_code == 404)
    r = c_vp.post(f"/api/admin/usuarios/{_uid}/sesiones/cerrar-todas", headers=_ha)
    chk("perfil ajeno: cerrar todas sus sesiones", r.status_code == 200 and r.json()["sesiones"] == [], r.text[:200])
    chk("perfil ajeno: su sesión ya no sirve", c_vp.get("/api/perfil", headers=_ho).status_code == 401)

# --- Observaciones del expediente ---
chk("observaciones: se conservan al limpiar", main.limpiar_datos_para_plantilla({"observaciones": "  Falta   firma\r\n\n\n\nRevisar folio "})["observaciones"] == "Falta firma\n\nRevisar folio")
with cliente() as c_ob:
    _h = {"Authorization": "Bearer " + c_ob.post("/api/auth/login", json={"username": f"ana_{SUF}", "password": "Prueba2026!"}).json()["token"]}
    c_ob.cookies.clear()
    _hadm = {"Authorization": "Bearer " + c_ob.post("/api/auth/login", json={"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123")}).json()["token"]}
    c_ob.cookies.clear()
    c_ob.post("/api/admin/usuarios", data={"username": f"obs_{SUF}", "password": "Prueba2026!", "rol": "operador"}, headers=_hadm)
    _ho = {"Authorization": "Bearer " + c_ob.post("/api/auth/login", json={"username": f"obs_{SUF}", "password": "Prueba2026!"}).json()["token"]}
    c_ob.cookies.clear()
    _cred = "17" + str(uuid.uuid4().int)[:8]
    r = c_ob.post("/api/expedientes/guardar", json={"expedientes": [{"ref": "o", "datos": {**DATOS, "numero_credito": _cred, "observaciones": "Nota al crear"}, "plantilla": PLANTILLA}]}, headers=_h)
    _id = r.json()["resultados"][0]["expediente_id"]
    _fila = lambda: next(x for x in c_ob.get("/api/expedientes", headers=_h).json() if x["id"] == _id)  # noqa: E731
    chk("observaciones: se guardan al crear", _fila()["observaciones"] == "Nota al crear", _fila().get("observaciones"))
    c_ob.put(f"/api/expedientes/{_id}/estado", json={"estado": "ENTREGADO"}, headers=_h)
    r = c_ob.put(f"/api/expedientes/{_id}/observaciones", json={"observaciones": "Cambiada"}, headers=_h)
    chk("observaciones: se cambian", r.status_code == 200 and r.json()["observaciones"] == "Cambiada", r.text[:200])
    chk("observaciones: no cambian el estado", _fila()["estado"] == "ENTREGADO")
    chk("observaciones: otro operador no puede", c_ob.put(f"/api/expedientes/{_id}/observaciones", json={"observaciones": "x"}, headers=_ho).status_code in (403, 404))
    r = c_ob.post("/api/expedientes/observaciones", json={"ids": [_id], "observaciones": "Agregada", "agregar": True}, headers=_h)
    chk("observaciones: agregar a varios", r.json()["observaciones"].get(_id) == "Cambiada\nAgregada", r.text[:200])
    r = c_ob.put(f"/api/expedientes/{_id}/editar", json={"datos": {**DATOS, "numero_credito": _cred, "observaciones": "Cambiada\nAgregada"}, "plantilla": PLANTILLA}, headers=_h)
    chk("observaciones: se conservan al editar", r.json()["expediente"]["observaciones"] == "Cambiada\nAgregada", r.text[:200])
    _bit = c_ob.get(f"/api/expedientes/{_id}/bitacora", headers=_h).json()["eventos"]
    chk("observaciones: quedan en el historial de cambios", any(e["tipo"] == "observaciones" for e in _bit))
    r = c_ob.post(f"/api/expedientes/{_id}/generar-word", json={"datos": {}, "plantilla": PLANTILLA}, headers=_h)
    chk("observaciones: no salen en el Word", r.status_code == 200 and "Agregada" not in texto_word(r.content))
    r = c_ob.post("/api/expedientes/generar-manual", json={**DATOS, "numero_credito": "18" + str(uuid.uuid4().int)[:8], "observaciones": "Desde captura", "plantilla": PLANTILLA}, headers=_h)
    chk("observaciones: captura manual", (r.json().get("datos_extraidos") or {}).get("observaciones") == "Desde captura", r.text[:300])

# --- Cada dato sale de su documento (el inmueble nunca de la carta) ---
_c = services.combinar_datos_pareja([{"_tipo_documento": "carta", "datos_inmueble": "AVENIDA VICENTE LOMBARDO TOLEDO VIV. 33", "numero_carta": "0543910959",
                                       "numero_credito": "1508007670", "fecha_expedicion": "08/09/2025", "folio_real": "999"}])
chk("solo carta: sus datos sí", _c["numero_carta"] == "0543910959" and _c["numero_credito"] == "1508007670" and _c["fecha_expedicion"] == "08/09/2025")
chk("solo carta: el inmueble y el folio no salen de la carta", _c["datos_inmueble"] == "NO_ENCONTRADO" and _c["folio_real"] == "NO_ENCONTRADO", _c)
_c = services.combinar_datos_pareja([{"_tipo_documento": "constancia", "fecha_expedicion": "01/01/2020", "numero_credito": "1508007670", "datos_inmueble": "LOTE 5"}])
chk("solo constancia: la fecha de expedición no sale de la constancia", _c["fecha_expedicion"] == "NO_ENCONTRADO" and _c["datos_inmueble"] == "LOTE 5", _c)
chk("solo constancia: el crédito sí (es el mismo dato)", _c["numero_credito"] == "1508007670")

print()
if FALLAS:
    print(f"HAY {len(FALLAS)} FALLA(S): {', '.join(FALLAS)}")
    sys.exit(1)
print("TODO OK")
